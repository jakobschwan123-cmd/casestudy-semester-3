# -*- coding: utf-8 -*-
"""Sprechtexte der Abschlusspräsentation. Schreibt tools/sprechtext.json (für die Notizen in pres.js)
und 4_praesentation/Sprechtext_Abschlusspraesentation.docx. Aufruf: python3 sprechtext.py"""
import json, os
from docx import Document
from docx.shared import Pt, Cm

HERE = os.path.dirname(os.path.abspath(__file__))

# (Folie, Titel, Vortragende(r), Sekunden, Sprechtext)
T = [
(1, 'Titel', 'Nina', 30,
 'Guten Morgen zusammen. Wir sind Gruppe 1 und stellen Ihnen RepairFlow vor: eine Plattform, mit der Reparaturwerkstätten ihre Aufträge, '
 'Ersatzteile und Kunden an einer Stelle steuern. Den Anfang macht David mit dem Problem, das wir lösen wollen. Danach zeigt Adrian die Lösung, '
 'Maximilian, Kilian und Jakob fassen zusammen, was wir dafür modelliert haben, und ich schließe mit Geschäftsmodell und Ausblick.'),
(2, 'Agenda', 'Nina', 40,
 'Unsere Präsentation hat zwei Hälften. In der ersten erzählen wir RepairFlow so, wie wir es einer Werkstatt vorstellen würden: Problem, Lösung, '
 'Nutzen, die Reparatur aus Kundensicht und unser Alleinstellungsmerkmal, die KI-Sofortdiagnose. In der zweiten Hälfte zeigen wir, was dahintersteckt: '
 'zehn Geschäftsprozesse, Rollen und Use Cases, das Datenmodell, unser Vorgehen und die Qualitätssicherung. Die Diagramme selbst finden Sie in der '
 'Projektdokumentation und im Abgabepaket; hier konzentrieren wir uns auf die Ergebnisse. Fragen gern am Ende, wir haben Zeit dafür eingeplant.'),
(3, 'Das Problem: Werkstätten heute', 'David', 75,
 'Stellen Sie sich eine Werkstattkette mit vier Filialen vor, die Fahrräder, E-Bikes und Elektronik repariert. In unserem Szenario heißt sie FixWerk. '
 'Der Alltag dort läuft über Papierzettel, Excel-Listen und Telefon. Das führt zu vier Problemen. Erstens: Der Kunde gibt sein Gerät ab und hört '
 'dann nichts mehr. Ob der Kostenvoranschlag freizugeben ist oder das Gerät fertig ist, erfährt er, wenn er anruft. Zweitens: Niemand weiß ohne '
 'Nachfrage, wo ein Auftrag gerade steht, weil Zettel und Listen auseinanderlaufen. Drittens: Jede Filiale bestellt Ersatzteile für sich, obwohl das '
 'Teil oft in der Nachbarfiliale liegt. Viertens: Termine werden nach Gefühl zugesagt, weil die Auslastung der Techniker nicht sichtbar ist. Die vier '
 'Punkte sind Annahmen unseres Szenarios, aber jeder, der schon einmal ein Gerät zur Reparatur gebracht hat, kennt mindestens eines davon.'),
(4, 'RepairFlow: die Lösung', 'Adrian', 70,
 'RepairFlow ist unsere Antwort darauf. Wichtig für die Einordnung: Wir sind nicht die Werkstatt, wir sind das Software-Startup. FixWerk ist unser '
 'Pilotkunde, an dessen Abläufen wir die Plattform entworfen haben. Diese Perspektive haben wir nach der Rückfrage des Dozenten bei der '
 'Konzeptvorstellung bewusst festgelegt. RepairFlow ist eine Cloud-Plattform, die den ganzen Reparatur-Lebenszyklus abdeckt, von der ersten Anfrage '
 'bis zur Abholung und Reklamation. Sie besteht aus vier Modulen: dem Kundenportal mit App, dem Auftrags- und Werkstattmodul, dem Dispositionsmodul '
 'für Ersatzteile und der Abrechnung. Zielkunden sind unabhängige Werkstätten und kleine Ketten mit zwei bis zehn Filialen. Außerhalb der '
 'Plattform bleiben die physische Reparatur, die Systeme der Lieferanten und die Buchhaltung; dorthin gibt es Schnittstellen, mehr nicht.'),
(5, 'Drei Nutzenversprechen', 'Adrian', 60,
 'Was hat die Werkstatt davon? Drei Dinge. Erstens: eine Plattform für alle drei Gerätearten, mit einer Kundenbasis und einem gemeinsamen '
 'Ersatzteillager. Zweitens: filialübergreifende Disposition. RepairFlow kennt die Bestände aller Standorte, reserviert dort, wo das Teil liegt, '
 'und bestellt nur, was nirgends vorhanden ist. Das spart Lieferzeit und Doppelbestellungen. Drittens: der Kunde bekommt Klarheit, bevor er die '
 'Werkstatt betritt. Mit einem Foto aus der App erhält er einen Diagnosevorschlag, einen vorläufigen Kostenvoranschlag und einen Termin. Für die '
 'Werkstatt heißt das weniger Telefon, weniger Standzeit im Lager und eine Auslastung, die sie wirklich sieht.'),
(6, 'Die Auftragsreise mit RepairFlow', 'David', 75,
 'So sieht eine Reparatur mit RepairFlow aus. Der Kunde schickt per App ein Foto oder Video und bekommt innerhalb einer Minute Vorschlag, '
 'vorläufigen Kostenvoranschlag und Terminvorschlag. Bestätigt er, sind die Teile schon vorreserviert. In der Filiale wird das Gerät registriert '
 'und der Auftrag angelegt. Der Techniker prüft den Vorschlag, stellt die Diagnose und erstellt den Kostenvoranschlag, den der Kunde wieder in der '
 'App freigibt. Dann disponiert RepairFlow die Teile über alle Filialen und bestellt Fehlteile beim Lieferanten. Nach der Reparatur folgt die '
 'Endkontrolle, der Kunde wird benachrichtigt, holt ab und zahlt. Jeder dieser Schritte ist ein Status des Reparaturauftrags: angenommen, in '
 'Diagnose, KVA offen, freigegeben, Teile bestellt, in Reparatur, fertig, abgeholt. Dieser Zustandsautomat ist der rote Faden durch alle unsere '
 'Modelle. Und falls doch etwas nicht passt: Die Reklamation startet einen eigenen Prozess, im Gewährleistungsfall kostenfrei.'),
(7, 'KI-Sofortdiagnose: KVA in 60 Sekunden', 'Adrian', 80,
 'Unser Alleinstellungsmerkmal ist die KI-Sofortdiagnose, die Sie gerade in der Auftragsreise gesehen haben. Fünf Schritte: Der Kunde nimmt Foto, '
 'Video oder Ton auf. Ein KI-Dienst erkennt Gerätetyp, wahrscheinliche Ursache und benötigte Teile. RepairFlow rechnet daraus einen vorläufigen '
 'Kostenvoranschlag, prüft die Verfügbarkeit in allen Filialen und reserviert vor. Zum Schluss schlägt es Filiale und Termin vor, und der Kunde '
 'bestätigt. Uns war wichtig, das glaubwürdig zu halten. Heute machbar ist die Erkennung sichtbarer Schäden per Bildanalyse, etwa ein gebrochenes '
 'Display oder eine gerissene Kette, dazu ein Vorab-KVA aus Teilepreisen und Aufwandsrichtwerten. Vision sind Geräuschanalyse bei E-Bike-Motoren '
 'und eine Aufwandsschätzung aus Video. Deshalb bleibt der Vorschlag immer vorläufig: Der Techniker bestätigt oder korrigiert ihn, und erst dann '
 'wird daraus ein verbindlicher Kostenvoranschlag. Die Medien werden nur mit Einwilligung verarbeitet und nach Abschluss gelöscht.'),
(8, 'Was wir gebaut haben: Ergebnisse in Zahlen', 'Maximilian', 50,
 'Damit zur zweiten Hälfte: Was steckt hinter dieser Geschichte? Wir haben drei Modelle erstellt, die zusammenpassen. Zehn BPMN-Kollaborationsdiagramme '
 'mit 133 Aktivitäten im Werkstatt-Pool, dazu die Abläufe des Kunden mit 26 weiteren Aktivitäten. Ein Use-Case-Diagramm mit 19 Use Cases und 17 '
 'Akteuren. Ein vollständiges Klassendiagramm mit 27 Klassen und 36 Assoziationen. Sechs Sequenzdiagramme, gefordert waren fünf, und als Zugabe ein '
 'Zustandsdiagramm für den Reparaturauftrag. Alle Modelle liegen im Camunda Modeler und in Visual Paradigm und werden mit der Dokumentation '
 'abgegeben. Ich beginne mit den Geschäftsprozessen.'),
(9, 'Zehn Prozesse, 61 Prozent automatisiert', 'Maximilian', 70,
 'Die zehn Prozesse decken den gesamten Lebenszyklus ab. Prozesse 01 bis 08 sind die durchgehende Auftragsreise: Sofortdiagnose, Annahme, Diagnose, '
 'Kostenvoranschlag, Ersatzteilreservierung, Bestellung, Reparatur und Abholung. 09 und 10 sind die Folgeprozesse Reklamation und Retoure. Jedes '
 'Diagramm ist eine Kollaboration: Der Werkstattbetrieb ist der ausführbare Pool mit den Lanes Service, Techniker, Werkstattleitung und Disposition, '
 'der Kunde hat seinen eigenen Ablauf, der Lieferant bleibt eine Black Box. Die wichtigste Zahl: 61 Prozent der Aktivitäten führt RepairFlow ohne '
 'manuelle Arbeit aus, also Statuswechsel, Benachrichtigungen, Verfügbarkeitsprüfungen, Berechnungen und Regeln. Was Handwerk ist, bleibt Handwerk: '
 'Diagnose, Reparatur und Endkontrolle sind Aktivitäten des Technikers mit Unterstützung durch das System.'),
(10, 'Drei Prozess-Highlights', 'Maximilian', 80,
 'Drei Stellen zeigen, was die Prozesse leisten. Erstens der Kostenvoranschlag in Prozess 04: Nach dem Versand wartet ein ereignisbasiertes Gateway '
 'auf Freigabe, Ablehnung oder den Ablauf von drei Tagen. Bei Fristablauf erinnert RepairFlow den Kunden, höchstens zweimal. Nach dem dritten '
 'Fristablauf ohne Antwort gilt der KVA als abgelehnt und das Gerät geht mit Diagnosepauschale zur Abholung. Kein Auftrag bleibt liegen. Zweitens die '
 'Disposition in Prozess 05: Für jedes benötigte Teil läuft parallel eine Prüfung, erst im eigenen Lager, dann in den anderen Filialen. Was lokal '
 'liegt, wird reserviert, was anderswo liegt, umgelagert, und nur Fehlteile werden bestellt. Drittens die Reklamation in Prozess 09: Der Techniker '
 'prüft den Mangel, RepairFlow entscheidet nach Regel über Gewährleistung, und im Gewährleistungsfall entsteht ein kostenfreier Nacharbeitsauftrag, '
 'der wieder durch die normale Auftragsreise läuft. So bleibt das Modell schlank, weil die Reklamation keine eigene Reparaturlogik braucht.'),
(11, 'Wer macht was im System: Rollen und Use Cases', 'David', 70,
 'Das Use-Case-Diagramm beschreibt, was RepairFlow für seine Benutzer leistet. 19 Use Cases in vier Bereichen: das Kundenportal mit Sofortdiagnose, '
 'Voranmeldung und KVA-Freigabe; Auftrag und Werkstatt von der Annahme bis zur Reklamation; die Ersatzteil-Disposition von der Reservierung bis zur '
 'Retoure; und die Administration des Werkstattbetriebs. Eine Besonderheit geht auf das Coaching zurück: Jeder Use Case hat genau eine Rolle. Statt '
 'eines Technikers mit fünf Verbindungen gibt es jetzt Diagnoseprüfer, Diagnostiker, Kalkulator, Reparateur und Endkontrolleur. Welche Person '
 'welche Rolle ausfüllt, regelt die Werkstatt selbst; im Datenmodell sind die Rollen den Mitarbeitertypen zugeordnet. Das Diagramm ist dadurch sofort '
 'lesbar, und die Rollen tauchen in den Sequenzdiagrammen wieder auf. Dazu kommen der Kunde und der KI-Diagnosedienst als externes System. '
 'Der Lieferant ist bewusst kein Akteur: Er hat keinen Zugang zu RepairFlow und bleibt eine Black Box, mit der nur Nachrichten ausgetauscht werden.'),
(12, 'Das Datenmodell: fünf Entwurfsentscheidungen', 'Kilian', 80,
 'Das Klassendiagramm ist ein einziges, vollständiges Modell mit 27 Klassen. Fünf Entscheidungen prägen es. Erstens: Der Reparaturauftrag ist die '
 'Aggregatwurzel. Befund und Reparaturschritte existieren nur mit ihrem Auftrag und sind deshalb Kompositionen, Kostenvoranschlag und Gerät sind nur '
 'zugeordnet, weil der Vorab-KVA schon vor dem Auftrag entsteht. Zweitens: Der Lagerbestand ist eine eigene Klasse zwischen Ersatzteil und Filiale. '
 'Nur so lässt sich die Verfügbarkeit über alle Filialen prüfen. Drittens: Die Reservierung ist eine eigene Klasse mit Lebenszyklus, vorreserviert, '
 'reserviert, entnommen, storniert. Viertens: Der Werkstattbetrieb ist der Mandant, und die Mitarbeiter sind nach Typ spezialisiert, '
 'Servicemitarbeiter, Techniker, Disponent und Werkstattleiter, passend zu den Lanes in BPMN und zu den Rollen im Use-Case-Diagramm. Fünftens: '
 'Voranmeldung, Medienanhang und KI-Diagnosevorschlag tragen die Sofortdiagnose. Die sechs Sequenzdiagramme verwenden ausschließlich Operationen '
 'dieser Klassen; jede Botschaft an ein Objekt ist eine Operation des Klassendiagramms.'),
(13, 'Vorgehen und Projektmanagement', 'Kilian', 70,
 'Gearbeitet haben wir in Sprints, die sich am Takt der vier Gruppentermine orientiert haben, mit Review, Retrospektive und Planning je Termin. '
 'Nina hat die Projektleitung, David die Stellvertretung und die Sicherungen, Adrian ist Product Owner, ich Scrum Master, Maximilian verantwortet '
 'BPMN, Jakob die Qualität. Alles Schriftliche liegt in einem GitHub-Repository; jede Änderung geht über einen eigenen Branch, einen Pull Request und '
 'ein Review durch eine zweite Person, bisher 27 gemergte Pull Requests. Das Trello-Board hat eine Liste je Termin und eine Karte je Person und '
 'Aufgabe. Jedes Diagramm wurde im Zweierteam geprüft. Die Meilensteine: Am 02.09. standen alle Diagramme im Entwurf, am 05.10. die Korrekturen aus '
 'zwei Gesamtprüfungen, am 15.10. haben wir die Rollenstruktur umgesetzt, am 22.10. war die Generalprobe. Die Abgabe folgt am 13.11.'),
(14, 'Qualität: Konventionen, Linter, Konsistenz', 'Jakob', 70,
 'Drei Modelle, die sich auf dieselben Dinge beziehen, laufen schnell auseinander. Dagegen hatten wir drei Mittel. Erstens Konventionen: '
 'Datenobjekte in BPMN tragen die Klassennamen, Aktivitäten heißen Verb plus Objekt, jede Rolle ist genau einer Lane und einer Klasse zugeordnet, und '
 'die Statuswerte sind in allen Modellen identisch. Zweitens Werkzeuge: Die BPMN-Dateien sind für Camunda 8 angereichert, mit Task-Definitionen, '
 'Timern und Bedingungen. Der Camunda-Linter und bpmnlint melden null Befunde. Drittens Prüfrunden: Eine Gesamtprüfung am 26.09., eine zweite Prüfrunde '
 'am 05.10. und der QA-Abgleich am 08.10. Jede Änderung wird gegen unsere QA-Checkliste geprüft, bevor sie gemergt wird. Ein '
 'Beispiel: Als wir die Nacharbeit in Prozess 07 auf zwei Durchläufe begrenzt haben, musste das Sequenzdiagramm SD4 nachziehen. Genau solche Stellen '
 'findet die Checkliste.'),
(15, 'Herausforderungen und Lernerfahrungen', 'Jakob', 70,
 'Vier Dinge haben uns am meisten beschäftigt. Die Perspektive: Werkstatt oder Softwareanbieter? Die Rückfrage des Dozenten hat uns zur klaren '
 'Entscheidung gezwungen, und danach mussten Texte und Klassenmodell nachziehen. Das Vokabular: Datenobjekte, Klassen, Lebenslinien und Statuswerte '
 'müssen wortgleich sein, sonst zerfällt die Analyse in drei Einzelmodelle. Die Werkzeuge: Der Camunda Modeler meldet für Camunda 8 fehlende '
 'technische Details als Fehler, und wir haben die Diagramme so lange angereichert, bis nichts mehr rot war. Und die Vision: Die KI-Sofortdiagnose '
 'soll beeindrucken, aber glaubwürdig bleiben, deshalb die Trennung in heute machbar und Vision. Gelernt haben wir, dass Konsistenz keine Prüfung am '
 'Ende ist, sondern eine Regel bei jeder Änderung. Und dass ein klarer Hinweis von außen, wie der zum Use-Case-Diagramm, ein Modell deutlich '
 'besser machen kann.'),
(16, 'Geschäftsmodell und Ausblick', 'Nina', 70,
 'Zum Schluss der Blick nach vorn. RepairFlow wird als Abo je Filiale angeboten, in zwei Stufen: Basis mit Auftrags- und Kundenverwaltung, Pro mit '
 'filialübergreifender Disposition und KI-Sofortdiagnose. Zielkunden sind unabhängige Werkstätten und kleine Ketten mit zwei bis zehn Filialen im '
 'deutschsprachigen Raum. Das sind Betriebe, die heute genau die vier Probleme vom Anfang haben und für die ein großes Werkstattsystem zu schwer '
 'ist. Der Ausblick geht in drei Richtungen: Predictive Disposition, also eine Bedarfsprognose je Filiale und Saison, damit Teile vor der '
 'Nachfrage im Lager liegen. Ein Techniker-Copilot, der Reparaturschritte aus der Historie vorschlägt. Und eine digitale Geräteakte, in der jedes '
 'Gerät seine Reparaturgeschichte behält. Nichts davon ist modelliert, aber alles baut auf dem Datenmodell auf, das wir heute vorgestellt haben.'),
(17, 'Fazit und Danke', 'Nina', 30,
 'RepairFlow löst den Medienbruch aus Papier, Excel und Telefon und macht den Reparaturauftrag zum zentralen Zustandsautomaten. Der fachliche '
 'Mehrwert liegt in der filialübergreifenden Disposition und in der KI-Sofortdiagnose, die dem Kunden Klarheit gibt, bevor er die Werkstatt betritt. '
 'Vielen Dank für Ihre Aufmerksamkeit. Wir freuen uns auf Ihre Fragen.'),
]

def main():
    json.dump([{'folie': f, 'titel': t, 'wer': w, 'sek': s, 'text': x} for f, t, w, s, x in T],
              open(os.path.join(HERE, 'sprechtext.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    total = sum(s for _, _, _, s, _ in T)
    d = Document()
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.5)
    p = d.add_paragraph(); r = p.add_run(f'Sprechtext Abschlusspräsentation RepairFlow (ca. {total // 60} Minuten)'); r.bold = True; r.font.size = Pt(14)
    d.add_paragraph('Gruppe 1, WWI25B4 · 27.10.2026, 09:00 Uhr, Raum B458 · Zeitangaben sind Richtwerte für die Generalprobe. '
                    'Der Text ist ein Vorschlag: Jede Person formuliert ihren Teil so um, dass er zu ihr passt.')
    per = {}
    for f, t, w, s, x in T:
        per[w] = per.get(w, 0) + s
    p = d.add_paragraph(); p.add_run('Redezeit je Person: ').bold = True
    p.add_run(', '.join(f'{w} {per[w] // 60}:{per[w] % 60:02d}' for w in ['Nina', 'David', 'Adrian', 'Maximilian', 'Kilian', 'Jakob']) + f' · gesamt {total // 60}:{total % 60:02d}.')
    for f, t, w, s, x in T:
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10)
        r = p.add_run(f'Folie {f} – {t}'); r.bold = True
        p.add_run(f'  ({s} s)').italic = True
        p = d.add_paragraph(); p.add_run(f'{w}: ').bold = True; p.add_run(x)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    p.add_run('Folie 18 – Anhang: Beiträge der Vortragenden').bold = True
    d.add_paragraph('Wird nicht vorgetragen. Pflichtangabe für den PDF-Export laut Ablauf.')
    out = os.path.join(HERE, '..', 'Sprechtext_Abschlusspraesentation.docx')
    d.save(out)
    print('ok', total, 's', per)

if __name__ == '__main__':
    main()
